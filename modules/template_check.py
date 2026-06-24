"""
模板残留 & 修订痕迹检测模块
① 模板占位符残留：【请填写】/ XXXX / （公司名称）等未替换内容
② Word 修订痕迹：未接受的修改、批注遗留（对方可以看到修改过程）
③ PDF 批注残留：便签、高亮、手写批注等
"""
import re
from dataclasses import dataclass, field
from typing import List, Optional

try:
    from docx import Document
    from docx.oxml.ns import qn
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

try:
    import fitz
    FITZ_OK = True
except ImportError:
    FITZ_OK = False


# ═══════════════════════════════════════════════════════════════
# 数据结构
# ═══════════════════════════════════════════════════════════════

@dataclass
class TemplateHit:
    hit_type: str       # placeholder / revision / annotation
    content: str        # 命中内容
    location: str       # 第N页 / 第N段
    risk_level: str     # 高/中
    suggestion: str


@dataclass
class TemplateCheckResult:
    placeholder_hits: List[TemplateHit] = field(default_factory=list)
    revision_hits: List[TemplateHit] = field(default_factory=list)
    annotation_hits: List[TemplateHit] = field(default_factory=list)
    problems: List[str] = field(default_factory=list)
    summary: str = ""
    file_type: str = ""


# ═══════════════════════════════════════════════════════════════
# ① 模板占位符检测（纯文本，适用所有文件类型）
# ═══════════════════════════════════════════════════════════════

# 常见占位符模式
PLACEHOLDER_PATTERNS = [
    # 中文方括号类
    (r'【[^】]{1,30}】',                "中文方括号占位符"),
    # 空白待填类
    (r'（?请(?:在此|此处)?填写[^）]{0,20}）?', "「请填写」未替换"),
    (r'请填[^\s，。]{0,15}',            "「请填…」未替换"),
    (r'（此处[^）]{1,20}）',            "「此处…」未替换"),
    # XXXX/___类
    (r'[Xx×_]{3,}',                    "XXXX/___占位符"),
    # 方括号英文类
    (r'\[(?:YOUR|COMPANY|INSERT|TODO|TBD|FILL)[^\]]{0,30}\]', "英文占位符"),
    # 年月日空白
    (r'\d{4}年\s*月\s*日',              "日期未填（年__月__日）"),
    (r'\d{4}年\d{1,2}月\s*日',         "日期未填（月__日）"),
    # 金额空白
    (r'人民币[（(]\s*[）)]\s*元',       "金额大写未填"),
    # 竞争对手/其他项目残留
    (r'(?:竞争对手|其他(?:公司|单位|投标人))[：:]\s*\S+', "疑似竞争对手信息残留"),
    # 上一个项目的名称残留（需要结合上下文，简单检测）
    (r'（以上内容[^）]{1,30}）',         "括号内提示文字未删除"),
    (r'注[：:]\s*[^。\n]{10,50}',       "模板注释未删除"),
]


def check_placeholders(text: str, pages: List[str] = None) -> List[TemplateHit]:
    hits: List[TemplateHit] = []
    covered_ranges = []   # 已覆盖的 (start, end) 区间，避免同一位置被多个模式重复命中

    def get_page(offset: int) -> str:
        if not pages:
            return "全文"
        cumlen = 0
        for i, pg in enumerate(pages):
            if cumlen + len(pg) >= offset:
                return f"第{i+1}页"
            cumlen += len(pg)
        return f"第{len(pages)}页"

    def overlaps(start: int, end: int) -> bool:
        for s, e in covered_ranges:
            if start < e and end > s:
                return True
        return False

    for pattern, label in PLACEHOLDER_PATTERNS:
        for m in re.finditer(pattern, text, re.IGNORECASE):
            val = m.group(0).strip()
            if len(val) < 2:
                continue
            if overlaps(m.start(), m.end()):
                continue   # 已被其他模式覆盖，跳过
            covered_ranges.append((m.start(), m.end()))
            loc = get_page(m.start())

            risk = "高" if any(kw in label for kw in ["占位符","未替换","未填"]) else "中"
            hits.append(TemplateHit(
                hit_type="placeholder",
                content=val[:60],
                location=loc,
                risk_level=risk,
                suggestion=f"将「{val[:20]}」替换为实际内容后再提交",
            ))

    return hits


# ═══════════════════════════════════════════════════════════════
# ② Word 修订痕迹检测
# ═══════════════════════════════════════════════════════════════

def check_word_revisions(file_bytes: bytes) -> List[TemplateHit]:
    """检测 Word 文档中的修订痕迹和批注"""
    if not DOCX_OK:
        return []
    hits: List[TemplateHit] = []
    try:
        import io
        doc = Document(io.BytesIO(file_bytes))
        body = doc.element.body

        # 检测修订插入/删除标记（w:ins / w:del）
        ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        ins_tags = body.findall(f'.//{{{ns}}}ins')
        del_tags = body.findall(f'.//{{{ns}}}del')

        if ins_tags:
            # 提取修订插入文字
            examples = []
            for tag in ins_tags[:3]:
                texts = [t.text for t in tag.iter() if hasattr(t,'tag') and 'r' in str(t.tag) and t.text]
                snippet = "".join(texts)[:30]
                if snippet:
                    examples.append(snippet)
            hits.append(TemplateHit(
                hit_type="revision",
                content=f"发现 {len(ins_tags)} 处修订插入" + (f"，如：「{'」「'.join(examples)}」" if examples else ""),
                location="Word文档",
                risk_level="高",
                suggestion="在 Word 中：审阅 → 接受 → 接受所有修订，清除修订痕迹后再另存提交",
            ))

        if del_tags:
            hits.append(TemplateHit(
                hit_type="revision",
                content=f"发现 {len(del_tags)} 处修订删除标记（被删内容对方可见）",
                location="Word文档",
                risk_level="高",
                suggestion="在 Word 中：审阅 → 接受 → 接受所有修订，或：审阅 → 显示标记 → 确认清除",
            ))

        # 检测批注（w:comment）
        comments_part = None
        try:
            comments_part = doc.part.comments_part
        except Exception:
            pass

        if comments_part is not None:
            comment_els = comments_part._element.findall(f'{{{ns}}}comment')
            if comment_els:
                # 提取批注内容示例
                examples = []
                for c in comment_els[:3]:
                    txt = "".join(t.text or "" for t in c.iter() if hasattr(t, 'text') and t.text)
                    if txt.strip():
                        examples.append(txt.strip()[:25])
                hits.append(TemplateHit(
                    hit_type="revision",
                    content=f"发现 {len(comment_els)} 条批注" + (f"，如：「{'」「'.join(examples)}」" if examples else ""),
                    location="Word文档",
                    risk_level="高",
                    suggestion="在 Word 中：审阅 → 删除 → 删除文档中的所有批注",
                ))

    except Exception as e:
        pass  # 文件损坏或格式不兼容时静默跳过

    return hits


# ═══════════════════════════════════════════════════════════════
# ③ PDF 批注残留检测
# ═══════════════════════════════════════════════════════════════

def check_pdf_annotations(file_bytes: bytes) -> List[TemplateHit]:
    """检测 PDF 中的批注、便签、高亮等"""
    if not FITZ_OK:
        return []
    hits: List[TemplateHit] = []
    try:
        import io
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        annot_pages = []
        total_annots = 0
        annot_types = set()

        for i, page in enumerate(doc):
            annots = list(page.annots())
            if annots:
                annot_pages.append(i + 1)
                total_annots += len(annots)
                for a in annots:
                    # fitz annot type names
                    try:
                        annot_types.add(a.type[1] if isinstance(a.type, (list, tuple)) else str(a.type))
                    except Exception:
                        pass

        if total_annots > 0:
            pages_str = "、".join(f"第{p}页" for p in annot_pages[:5])
            if len(annot_pages) > 5:
                pages_str += f"等{len(annot_pages)}页"
            types_str = "、".join(str(t) for t in list(annot_types)[:4]) if annot_types else "批注"
            hits.append(TemplateHit(
                hit_type="annotation",
                content=f"PDF 含 {total_annots} 个批注/标注（{types_str}），位于{pages_str}",
                location=pages_str,
                risk_level="中",
                suggestion="用 Adobe Acrobat：编辑 → 管理批注 → 全选 → 删除；或用 PDF 打印另存为新文件",
            ))
    except Exception:
        pass
    return hits


# ═══════════════════════════════════════════════════════════════
# 主入口
# ═══════════════════════════════════════════════════════════════

def check_template_and_revisions(
    file_bytes: bytes,
    filename: str,
    text: str,
    pages: List[str] = None,
) -> TemplateCheckResult:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    result = TemplateCheckResult(file_type=ext)

    # ① 占位符（所有格式）
    if text.strip():
        result.placeholder_hits = check_placeholders(text, pages)

    # ② Word 修订痕迹
    if ext in ("docx", "doc") and file_bytes and DOCX_OK:
        result.revision_hits = check_word_revisions(file_bytes)

    # ③ PDF 批注
    if ext == "pdf" and file_bytes and FITZ_OK:
        result.annotation_hits = check_pdf_annotations(file_bytes)

    # 汇总 problems
    all_hits = result.placeholder_hits + result.revision_hits + result.annotation_hits
    for h in all_hits:
        icon = "🔴" if h.risk_level == "高" else "🟠"
        result.problems.append(f"{icon} {h.location}：{h.content[:60]}")

    # summary
    ph = len(result.placeholder_hits)
    rv = len(result.revision_hits)
    an = len(result.annotation_hits)
    total = ph + rv + an

    if total == 0:
        result.summary = "✅ 未发现模板残留、修订痕迹或批注"
    else:
        parts = []
        if ph: parts.append(f"模板占位符 {ph} 处")
        if rv: parts.append(f"修订痕迹/批注 {rv} 项")
        if an: parts.append(f"PDF批注 {an} 处")
        result.summary = f"⚠️ 发现：{'、'.join(parts)} — 提交前必须清除！"

    return result
